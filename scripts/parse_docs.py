import os
import re
import json
import sys
import docx

def clean_text(text):
    if not text:
        return ""
    text = text.replace('\xa0', ' ').replace('\u2018', "'").replace('\u2019', "'")
    text = text.replace('\u201c', '"').replace('\u201d', '"')
    return text.strip()

def normalize_number_str(s):
    return re.sub(r'(\d+),(\d+)', r'\1.\2', s.strip())

def clean_header_name(h):
    h = h.replace('φ', 'phi').replace('Φ', 'phi').replace('π', 'pi')
    clean_h = re.sub(r'[^a-zA-Z0-9]', '_', h)
    clean_h = re.sub(r'_+', '_', clean_h).strip('_')
    if not clean_h:
        clean_h = 'param'
    return clean_h

def parse_split_variant_table(table, header_pattern):
    rows = table.rows
    if len(rows) < 2:
        return []
    
    header_cells = [clean_text(c.text) for c in rows[0].cells]
    half = len(header_cells) // 2
    left_headers = header_cells[:half]
    right_headers = header_cells[half:]
    
    variants = []
    
    for r in rows[1:]:
        cells = [clean_text(c.text) for c in r.cells]
        if not any(cells):
            continue
        # Left half
        left_c = cells[:half]
        if left_c and left_c[0].isdigit():
            var_num = int(left_c[0])
            params = {}
            for h, v in zip(left_headers[1:], left_c[1:]):
                clean_h = clean_header_name(h)
                params[clean_h] = normalize_number_str(v)
            variants.append({"variant": var_num, "parameters": params})
            
        # Right half
        right_c = cells[half:]
        if right_c and right_c[0].isdigit():
            var_num = int(right_c[0])
            params = {}
            for h, v in zip(right_headers[1:], right_c[1:]):
                clean_h = clean_header_name(h)
                params[clean_h] = normalize_number_str(v)
            variants.append({"variant": var_num, "parameters": params})
            
    variants.sort(key=lambda x: x["variant"])
    return variants

def parse_single_variant_table(table):
    rows = table.rows
    if len(rows) < 2:
        return []
    
    header_cells = [clean_text(c.text) for c in rows[0].cells]
    variants = []
    
    for r in rows[1:]:
        cells = [clean_text(c.text) for c in r.cells]
        if not any(cells) or not cells[0].isdigit():
            continue
        var_num = int(cells[0])
        params = {}
        for h, v in zip(header_cells[1:], cells[1:]):
            clean_h = clean_header_name(h)
            params[clean_h] = normalize_number_str(v)
        variants.append({"variant": var_num, "parameters": params})
        
    variants.sort(key=lambda x: x["variant"])
    return variants

def parse_docx_lab(file_path, lab_number):
    doc = docx.Document(file_path)
    paragraphs = [clean_text(p.text) for p in doc.paragraphs if clean_text(p.text)]
    
    topic = ""
    for idx, p in enumerate(paragraphs[:5]):
        if "mavzu:" in p.lower():
            topic = p.split(":", 1)[1].strip()
            if idx + 1 < len(paragraphs) and not any(k in paragraphs[idx+1].lower() for k in ['ishning maqsadi', 'maqsad', 'nazariy']):
                topic += " " + paragraphs[idx+1]
            break
        elif "mashg'ulot" in p.lower() or "mashg‘ulot" in p.lower():
            if idx + 1 < len(paragraphs) and "ishning maqsadi" not in paragraphs[idx+1].lower():
                topic = paragraphs[idx+1]
                
    if not topic and len(paragraphs) > 1:
        topic = paragraphs[1]
    topic = clean_text(topic.strip('. '))
    
    objective = ""
    for idx, p in enumerate(paragraphs):
        if "ishning maqsadi:" in p.lower() or "maqsad:" in p.lower():
            objective = p.split(":", 1)[1].strip()
            j = idx + 1
            while j < len(paragraphs) and not any(k in paragraphs[j].lower() for k in ['nazariy qism', 'topshiriq', '1-topshiriq', 'bajarish tartibi']):
                objective += " " + paragraphs[j]
                j += 1
            break
            
    theory_paras = []
    tasks_paras = []
    report_paras = []
    questions_paras = []
    conclusion_paras = []
    code_snippets = []
    
    current_section = "theory"
    code_keywords = ['subplot(', 'figure;', 'plot(', 'stem(', 'imread(', 'imshow(', 'histeq(', 'fft(', 'conv(', 'filter(', 'dct2(', 'medfilt2(']
    current_code = []
    
    for idx, p in enumerate(paragraphs):
        p_lower = p.lower()
        if "nazariy qism" in p_lower:
            current_section = "theory"
            continue
        elif any(k in p_lower for k in ['topshiriq variantlari:', 'topshiriq:', 'mustaqil topshiriqlar', 'bajaradigan topshiriqlari', 'amaliy ish bo‘yicha topshiriq', 'amaliy ish bo\'yicha topshiriq']):
            current_section = "tasks"
        elif "hisobot tarkibi" in p_lower or "hisobot mazmuni" in p_lower:
            current_section = "report"
            continue
        elif "nazorat savollari" in p_lower:
            current_section = "questions"
            continue
        elif "xulosa" in p_lower and len(p) < 40 and idx > len(paragraphs) // 2:
            current_section = "conclusion"
            continue
            
        is_code = any(k in p for k in code_keywords) or (p.startswith('%') and len(p) > 2) or (p.startswith(('Fs =', 'fs =', 't =', 'x =', 'y =', 'I =', 'z =', 'f =', 'A =', 'N =', 'n =')))
        if is_code or (current_code and (p.endswith(';') or p.endswith('end') or p.endswith(')'))):
            current_code.append(p)
        else:
            if current_code:
                code_snippets.append("\n".join(current_code))
                current_code = []
                
        if current_section == "theory":
            theory_paras.append(p)
        elif current_section == "tasks":
            tasks_paras.append(p)
        elif current_section == "report":
            report_paras.append(p)
        elif current_section == "questions":
            questions_paras.append(p)
        elif current_section == "conclusion":
            conclusion_paras.append(p)
            
    if current_code:
        code_snippets.append("\n".join(current_code))
        
    questions = []
    for q in questions_paras:
        q_clean = q.strip()
        q_norm = re.sub(r'^\d+[\.\)\-]\s*', '', q_clean)
        if q_norm.endswith('?') or len(q_norm) > 15:
            questions.append(q_norm)
            
    variant_tables = []
    for t_idx, table in enumerate(doc.tables):
        t_data = []
        for row in table.rows:
            t_data.append([clean_text(cell.text) for cell in row.cells])
            
        header_row = " ".join([c.lower() for c in t_data[0]]) if t_data else ""
        if "№" in header_row or "variant" in header_row or "f1" in header_row or "x ( n )" in header_row or "chastota" in header_row:
            if len(table.columns) >= 4 and len([c for c in t_data[0] if "№" in c or "variant" in c.lower()]) >= 2:
                vars_list = parse_split_variant_table(table, header_row)
                if vars_list:
                    variant_tables.append({"table_index": t_idx, "type": "split_20", "variants": vars_list})
            else:
                vars_list = parse_single_variant_table(table)
                if vars_list:
                    variant_tables.append({"table_index": t_idx, "type": "single", "variants": vars_list})
                    
    final_variants = []
    if variant_tables:
        final_variants = variant_tables[0]["variants"]
        if lab_number == 14 and len(final_variants) == 10:
            for n in range(11, 21):
                base_var = final_variants[(n - 1) % 10]
                final_variants.append({
                    "variant": n,
                    "parameters": dict(base_var["parameters"]),
                    "mapped_from": (n - 1) % 10 + 1
                })
    else:
        final_variants = generate_default_variants_for_lab(lab_number)
        
    return {
        "lab_number": lab_number,
        "title": topic,
        "objective": objective,
        "theory": "\n\n".join(theory_paras),
        "tasks": "\n\n".join(tasks_paras),
        "report_requirements": "\n\n".join(report_paras),
        "conclusion": "\n\n".join(conclusion_paras),
        "code_snippets": code_snippets,
        "control_questions_raw": questions,
        "variants": final_variants,
        "raw_tables_count": len(doc.tables)
    }

def generate_default_variants_for_lab(lab_num):
    variants = []
    for v in range(1, 21):
        params = {}
        if lab_num == 8:
            noise_types = ["salt_pepper", "gaussian", "speckle"]
            n_type = noise_types[(v - 1) % 3]
            var_val = round(0.02 + (v % 5) * 0.02, 3)
            params = {"noise_type": n_type, "noise_param": var_val, "filter_size": 3 if v % 2 == 1 else 5}
        elif lab_num == 9:
            shapes = ["disk", "square", "diamond", "line"]
            params = {"se_shape": shapes[(v - 1) % 4], "se_radius": 2 + (v % 6) * 2, "operation": ["open", "close", "gradient", "tophat"][(v-1)%4]}
        elif lab_num == 11:
            q_values = [10, 20, 30, 40, 50, 60, 70, 80, 90]
            params = {"Q": q_values[(v - 1) % len(q_values)], "keep_coeffs": 4 + (v % 12) * 2}
        elif lab_num == 12:
            params = {"quality": 30 + (v % 7) * 10, "target_format": "jpeg"}
        elif lab_num == 13:
            params = {"window_ms": 20 + (v % 3) * 10, "overlap_pct": 50, "lpc_order": 10 + (v % 5)}
        elif lab_num == 15:
            params = {"iou_threshold": round(0.3 + (v % 5) * 0.1, 2), "score_threshold": round(0.4 + (v % 4) * 0.1, 2)}
        elif lab_num == 16:
            params = {"upscale_factor": 2 if v % 2 == 1 else 4, "sample_rate_sparse": 2 + (v % 4)}
        elif lab_num == 17:
            params = {"target_model": ["HSV", "YCbCr", "Grayscale"][(v - 1) % 3], "target_channel": ["H", "S", "V", "Y", "Cb", "Cr"][(v - 1) % 6]}
        elif lab_num == 18:
            params = {"k_neighbors": 1 if v % 3 == 1 else (3 if v % 3 == 2 else 5), "train_split": 0.6 + (v % 3) * 0.1}
        elif lab_num == 19:
            params = {"cutoff_hz": 2000 + (v % 5) * 500, "filter_order": 4 + (v % 3) * 2}
        elif lab_num == 20:
            params = {"bp_low_hz": 200 + (v % 5) * 50, "bp_high_hz": 2500 + (v % 5) * 200, "sharpen_radius": 1 + (v % 3)}
        elif lab_num == 21:
            params = {"motion_len": 10 + (v % 5) * 3, "motion_theta": 15 + (v % 6) * 15, "nsr": round(0.005 + (v % 5) * 0.005, 4)}
        else:
            params = {"param_a": v * 10, "param_b": v * 5}
            
        variants.append({"variant": v, "parameters": params})
    return variants

if __name__ == "__main__":
    docs_dir = r"C:\Users\User\.gemini\antigravity\scratch\signallab\source_docs"
    out_dir = r"C:\Users\User\.gemini\antigravity\scratch\signallab\content\labs"
    
    print("Re-parsing all 20 docx files with clean parameter names...")
    for num in range(2, 22):
        fname = f"{num}-Amaliy mashg'ulot.docx"
        fpath = os.path.join(docs_dir, fname)
        if not os.path.exists(fpath):
            fname = f"{num}-Amaliy_mashg_ulot.docx"
            fpath = os.path.join(docs_dir, fname)
            
        if not os.path.exists(fpath):
            continue
            
        lab_data = parse_docx_lab(fpath, num)
        out_path = os.path.join(out_dir, f"{num}.json")
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(lab_data, f, ensure_ascii=False, indent=2)
            
    print("Completed re-parsing!")
